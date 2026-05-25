"""Commission Invoice Excel export.

Self-contained module — no imports from purchase_orders.py or po_exports.py.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.commission_invoices import _float, safe_ci_filename
from app.database import get_data_dir

BASE       = Path(__file__).resolve().parent.parent
XLSX_DIR   = get_data_dir() / "exports" / "commission_invoices" / "xlsx"
LOGO_PNG   = BASE / "static" / "gbbv-logo.png"
FOOTER_PNG = BASE / "static" / "footer logo.png"

# ── Style constants ──────────────────────────────────────────────────────────
_GREEN  = "1A5632"
_WHITE  = "FFFFFF"
_LGREY  = "F2F2F2"
_BORDER = "B0C4B8"

_thin  = Side(style="thin",   color=_BORDER)
_thick = Side(style="medium", color="888888")

_all_thin   = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_all_thick  = Border(left=_thick, right=_thick, top=_thick, bottom=_thick)
_bottom_med = Border(bottom=Side(style="medium", color=_GREEN))

_hdr_font  = Font(bold=True, color=_WHITE, size=9)
_bold9     = Font(bold=True, size=9)
_reg9      = Font(size=9)
_hdr_fill  = PatternFill("solid", fgColor=_GREEN)
_grey_fill = PatternFill("solid", fgColor=_LGREY)
_c         = Alignment(horizontal="center", vertical="center", wrap_text=True)
_l         = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_r         = Alignment(horizontal="right",  vertical="center")


def _s(ws, coord_or_cell, value=None, *, bold=False, fill=None, border=None,
        align=None, font=None):
    """Set value + style on a cell (coord string or cell object)."""
    cell = ws[coord_or_cell] if isinstance(coord_or_cell, str) else coord_or_cell
    if value is not None:
        cell.value = value
    if font:
        cell.font = font
    elif bold:
        cell.font = _bold9
    else:
        cell.font = _reg9
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if align:
        cell.alignment = align
    return cell


def _add_logo(ws) -> None:
    if not LOGO_PNG.exists():
        return
    try:
        img = XLImage(str(LOGO_PNG))
        img.width  = 80
        img.height = 60
        ws.add_image(img, "A1")
    except Exception:
        pass


def _money(val: Any) -> str:
    v = _float(val)
    return f"{v:,.2f}" if v else "0.00"


def build_ci_workbook(ci: dict[str, Any]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commission Invoice"

    # ── column widths ───────────────────────────────────────────────────────
    widths = [20, 16, 10, 12, 14, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Logo placeholder (rows 1-3) ─────────────────────────────────────────
    _add_logo(ws)
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20

    # ── Company name (B4) ───────────────────────────────────────────────────
    ws.row_dimensions[4].height = 16
    _s(ws, "B4", ci.get("company_name") or "Godavari Biorefineries Inc",
       bold=True, align=_l)

    # ── Document title (A9) ─────────────────────────────────────────────────
    for r in range(5, 9):
        ws.row_dimensions[r].height = 12
    ws.merge_cells("A9:D9")
    _s(ws, "A9", ci.get("document_title") or "COMMISSION INVOICE",
       font=Font(bold=True, size=13, color=_GREEN), align=_c,
       border=_bottom_med)
    ws.row_dimensions[9].height = 20

    # ── Bill-to + Invoice meta (rows 13-18) ────────────────────────────────
    ws.row_dimensions[13].height = 13
    _s(ws, "A13", "TO", bold=True, align=_l)
    _s(ws, "D13", "Invoice No", bold=True, align=_r)
    _s(ws, "E13", ci.get("invoice_number") or "", align=_l)

    _s(ws, "A14", ci.get("bill_to_name") or "", bold=True, align=_l)
    _s(ws, "D14", "Date", bold=True, align=_r)
    _s(ws, "E14", ci.get("invoice_date") or "", align=_l)

    for r, fld in [(15, "bill_to_address_1"), (16, "bill_to_address_2"), (17, "bill_to_address_3")]:
        ws.row_dimensions[r].height = 13
        _s(ws, f"A{r}", ci.get(fld) or "", align=_l)

    ws.row_dimensions[18].height = 13
    _s(ws, "A18", "Customer Order no", bold=True, align=_l)
    _s(ws, "C18", ci.get("customer_order_no") or "", align=_l)

    # ── Transaction description (rows 20-21) ───────────────────────────────
    for r in (19, 20, 21):
        ws.row_dimensions[r].height = 13
    ws.merge_cells("A20:G20")
    _s(ws, "A20", "TRANSACTION DESCRIPTION", bold=True, fill=_grey_fill,
       border=_all_thin, align=_c)
    ws.merge_cells("A21:G21")
    _s(ws, "A21", ci.get("transaction_description") or "", align=_l, border=_all_thin)
    ws.row_dimensions[21].height = 28

    # ── Shipment details (rows 22-26) ──────────────────────────────────────
    for r in range(22, 27):
        ws.row_dimensions[r].height = 13

    ws.merge_cells("A22:D22")
    _s(ws, "A22", f"Shipment Direct from GBL on  {ci.get('shipment_date') or ''}", align=_l)
    _s(ws, "E22", "BL No", bold=True, align=_r)
    ws.merge_cells("F22:G22")
    _s(ws, "F22", ci.get("bl_number") or "", align=_l)

    ws.merge_cells("A23:D23")
    _s(ws, "A23", f"Port of Loading : {ci.get('port_of_loading') or ''}", align=_l)
    _s(ws, "E23", "BL Date", bold=True, align=_r)
    ws.merge_cells("F23:G23")
    _s(ws, "F23", ci.get("bl_date") or "", align=_l)

    ws.merge_cells("A24:G24")
    _s(ws, "A24", f"Container No :  {ci.get('container_numbers') or ''}", align=_l)

    # ── Product table header (rows 26-27) ──────────────────────────────────
    ws.row_dimensions[26].height = 16
    ws.row_dimensions[27].height = 16

    headers_top = [
        ("A26", "C26", "PRODUCT AND OTHER DESCRIPTION"),
        ("D26", "D26", "GBL Invoice #"),
        ("E26", "E26", "QUANTITY"),
        ("F26", "F26", "Value"),
        ("G26", "G26", "Commission"),
    ]
    for start, end, txt in headers_top:
        if start != end:
            ws.merge_cells(f"{start}:{end}")
        _s(ws, start, txt, font=_hdr_font, fill=_hdr_fill, border=_all_thin, align=_c)

    headers_sub = [
        ("A27", "C27", ""),
        ("D27", "D27", ""),
        ("E27", "E27", "MT"),
        ("F27", "F27", "FOB Euro"),
        ("G27", "G27", "Euro"),
    ]
    for start, end, txt in headers_sub:
        if start != end:
            ws.merge_cells(f"{start}:{end}")
        _s(ws, start, txt, font=_hdr_font, fill=_hdr_fill, border=_all_thin, align=_c)

    # ── Line items (from row 28) ────────────────────────────────────────────
    r = 28
    lines = ci.get("line_items") or []
    for li in lines:
        ws.row_dimensions[r].height = 14
        ws.merge_cells(f"A{r}:C{r}")
        _s(ws, f"A{r}", li.get("product_description") or "", align=_l, border=_all_thin)
        _s(ws, f"D{r}", li.get("gbl_invoice_number") or "",  align=_l, border=_all_thin)
        _s(ws, f"E{r}", _float(li.get("quantity")),           align=_r, border=_all_thin)
        _s(ws, f"F{r}", _float(li.get("fob_value")),          align=_r, border=_all_thin)

        # Commission rate + value in same cell group
        rate_txt = f"{_float(li.get('commission_rate')):g}% on FOB"
        _s(ws, f"G{r}", rate_txt, align=_c, border=_all_thin)
        r += 1

        ws.row_dimensions[r].height = 13
        ws.merge_cells(f"A{r}:F{r}")
        _s(ws, f"A{r}", "", border=_all_thin)
        cv = _float(li.get("commission_value"))
        _s(ws, f"G{r}", cv, align=_r, border=_all_thin)
        r += 1

    # ── Totals block ────────────────────────────────────────────────────────
    r += 1
    ws.row_dimensions[r].height = 13
    ws.merge_cells(f"A{r}:F{r}")
    _s(ws, f"A{r}", "TOTAL", bold=True, fill=_grey_fill, border=_all_thin, align=_r)
    _s(ws, f"G{r}", ci.get("total_commission") or 0, bold=True,
       fill=_grey_fill, border=_all_thin, align=_r)
    r += 1

    for label, key in [("Net VALUE", "net_value"), ("VAT (0 %)", "vat_amount"),
                        ("TOTAL TO PAY", "total_to_pay")]:
        ws.row_dimensions[r].height = 13
        ws.merge_cells(f"A{r}:D{r}")
        _s(ws, f"A{r}", label, bold=(label == "TOTAL TO PAY"),
           fill=_grey_fill if label == "TOTAL TO PAY" else None,
           border=_all_thin, align=_l)
        ws.merge_cells(f"E{r}:G{r}")
        _s(ws, f"E{r}", ci.get(key) or 0,
           bold=(label == "TOTAL TO PAY"),
           fill=_grey_fill if label == "TOTAL TO PAY" else None,
           border=_all_thin, align=_r)
        r += 1

    # ── In Words ────────────────────────────────────────────────────────────
    r += 1
    ws.row_dimensions[r].height = 13
    ws.merge_cells(f"A{r}:G{r}")
    _s(ws, f"A{r}", f"In Words :  {ci.get('amount_in_words') or ''}", align=_l, border=_all_thin)

    # ── Payment Terms ───────────────────────────────────────────────────────
    r += 2
    ws.row_dimensions[r].height = 13
    _s(ws, f"A{r}", "Payment Terms", bold=True, align=_l)
    ws.merge_cells(f"C{r}:G{r}")
    _s(ws, f"C{r}", ci.get("payment_terms") or "Prompt", align=_l)

    # ── Enclosures ──────────────────────────────────────────────────────────
    r += 1
    ws.row_dimensions[r].height = 13
    _s(ws, f"A{r}", "Enclosures", bold=True, align=_l)
    ws.merge_cells(f"C{r}:G{r}")
    _s(ws, f"C{r}", ci.get("enclosures") or "", align=_l)

    # ── Bank Details ────────────────────────────────────────────────────────
    r += 2
    ws.row_dimensions[r].height = 13
    _s(ws, f"A{r}", "BANK DETAILS", bold=True, fill=_grey_fill, align=_l, border=_all_thin)
    r += 1
    ws.merge_cells(f"C{r}:G{r}")
    _s(ws, f"C{r}", ci.get("bank_name") or "", bold=True, align=_l)
    r += 1
    ws.merge_cells(f"C{r}:G{r}")
    _s(ws, f"C{r}", ci.get("bank_account_no") or "", align=_l)
    r += 1
    _s(ws, f"C{r}", "IBAN", bold=True, align=_l)
    _s(ws, f"D{r}", ci.get("bank_iban") or "", align=_l)
    ws.merge_cells(f"F{r}:G{r}")
    _s(ws, f"F{r}", f"BIC : {ci.get('bank_bic') or ''}", align=_l)

    # ── Signature block ─────────────────────────────────────────────────────
    r += 2
    ws.merge_cells(f"A{r}:C{r}")
    _s(ws, f"A{r}", "For Godavari Biorefineries Inc", bold=True, align=_l)
    r += 4
    ws.merge_cells(f"A{r}:C{r}")
    _s(ws, f"A{r}", "Authorised Signatory", bold=True,
       border=Border(top=Side(style="medium", color="333333")), align=_c)

    # ── Footer ──────────────────────────────────────────────────────────────
    r += 3
    ws.merge_cells(f"C{r}:D{r}")
    _s(ws, f"C{r}", "Commercial Register Number", align=_r)
    _s(ws, f"E{r}", 34325188, align=_l)
    ws.merge_cells(f"F{r}:G{r}")
    _s(ws, f"F{r}", "VAT : NL8203.86.157.B.01", align=_l)
    r += 1
    ws.merge_cells(f"C{r}:E{r}")
    _s(ws, f"C{r}",
       "Godavari Biorefineries Inc   Opaallaan 1180,   2132 LN,   Hoofddorp", align=_l)
    r += 1
    ws.merge_cells(f"C{r}:D{r}")
    _s(ws, f"C{r}", "The Netherlands", align=_l)
    ws.merge_cells(f"F{r}:G{r}")
    _s(ws, f"F{r}", "TEL : + 31 6 11 12 61 66", align=_l)

    return wb


def export_ci_xlsx(ci: dict[str, Any]) -> tuple[bytes, str]:
    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    fname = f"CI_{safe_ci_filename(ci.get('invoice_number', 'CI'))}.xlsx"
    path  = XLSX_DIR / fname
    wb    = build_ci_workbook(ci)
    wb.save(path)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read(), fname
